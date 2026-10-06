import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
from datetime import datetime

# Shrift va ranglar (De Luxury premium uslubida)
HEADER_FILL = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid") # To'q ko'k
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="1A365D")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="4A5568")
TOTAL_FILL = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color="000000")
REGULAR_FONT = Font(name="Calibri", size=10)

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E0"),
    right=Side(style="thin", color="CBD5E0"),
    top=Side(style="thin", color="CBD5E0"),
    bottom=Side(style="thin", color="CBD5E0")
)

def adjust_columns_width(ws, start_row: int = 4, min_width: int = 12):
    """Sarlavhalarni hisobga olmasdan faqat jadval ustunlari bo'yicha chiroyli kenglik o'rnatadi"""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col_letter == 'A':
            ws.column_dimensions[col_letter].width = 6
            continue
        elif col_letter == 'B':
            ws.column_dimensions[col_letter].width = 13
            continue

        # Faqat 4-qatordan pastdagi katakchalarni o'lchaymiz
        max_len = 0
        for cell in col:
            if cell.row >= start_row and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, min_width)

class ExcelGenerator:
    @staticmethod
    def generate_stock_excel(stock_data: dict, folder_name: str = None, output_dir: str = "temp") -> str:
        """Qoldiqlar bo'yicha to'liq (yoki guruh bo'yicha filtrlangan) Excel hisoboti yaratadi"""
        os.makedirs(output_dir, exist_ok=True)
        folder_suffix = f"_{folder_name}" if folder_name else ""
        clean_suffix = "".join(c for c in folder_suffix if c.isalnum() or c in (' ', '_', '-')).strip()
        filename = f"Qoldiqlar_{clean_suffix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx" if clean_suffix else f"Qoldiqlar_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(output_dir, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Qoldiqlar"
        ws.views.sheetView[0].showGridLines = True

        # Sarlavhani 1-2 qatorlarda birlashtirib chiroyli formatlaymiz
        ws.merge_cells("A1:K2")
        title_cell = ws["A1"]
        title_text = f"DE LUXURY — TOVAR QOLDIQLARI ({folder_name.upper()} GURUHI)" if folder_name else "DE LUXURY — TOVAR QOLDIQLARI VA BAHOLASH HISOBOTI"
        title_cell.value = title_text
        title_cell.font = TITLE_FONT
        title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("A3:K3")
        sub_cell = ws["A3"]
        sub_cell.value = f"Shakllantirilgan sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        sub_cell.font = SUBTITLE_FONT
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")

        headers = [
            "№", "Tovar kodi", "Tovar nomi", "Guruhi", "Kategoriya", "Vid tovara",
            "Qoldiq (dona)", "Tan narxi (so'm)", "Jami tan narxi", "Sotuv narxi (so'm)", "Jami sotuv narxi"
        ]

        row_num = 4
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        raw_rows = stock_data.get("raw_rows", [])
        if folder_name:
            raw_rows = [r for r in raw_rows if r.get("folder", "").lower() == folder_name.lower()]

        idx = 1
        row_num = 5

        total_qty = 0.0
        total_cost = 0.0
        total_sale = 0.0

        for item in raw_rows:
            qty = float(item.get("stock", 0) or 0)
            if qty <= 0:
                continue

            code = item.get("code", "")
            name = item.get("name", "")
            folder = item.get("folder", "Guruhsiz")
            cat = item.get("category", "")
            vid = item.get("vid_tovara", "")
            
            buy_price = float(item.get("buy_price", 0) or 0)
            cost_total = float(item.get("cost_total", 0) or 0)
            sale_price = float(item.get("sale_price", 0) or 0)
            sale_total = float(item.get("sale_total", 0) or 0)

            total_qty += qty
            total_cost += cost_total
            total_sale += sale_total

            values = [
                idx, code, name, folder, cat, vid,
                qty, buy_price, cost_total, sale_price, sale_total
            ]

            for col_num, val in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.font = REGULAR_FONT
                cell.border = THIN_BORDER
                
                # Tekislash va son formatlari
                if col_num in [1, 2]: # №, Kod
                    cell.alignment = Alignment(horizontal="center")
                elif col_num in [3, 4, 5, 6]: # Nomi, Guruh, Kategoriya, Vid
                    cell.alignment = Alignment(horizontal="left")
                elif col_num == 7: # Qoldiq
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0"
                elif col_num in [8, 9, 10, 11]: # Narxlar
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0.00"

            row_num += 1
            idx += 1

        # Jami qatori
        total_cells = [
            "JAMI", "", "", "", "", "",
            total_qty, "", total_cost, "", total_sale
        ]
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=6)
        for col_num, val in enumerate(total_cells, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val if val != "" else None)
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = THIN_BORDER
            if col_num in [7, 9, 11]:
                cell.alignment = Alignment(horizontal="right")
                if col_num == 7:
                    cell.number_format = "#,##0"
                else:
                    cell.number_format = "#,##0.00"
        
        ws.cell(row=row_num, column=1).alignment = Alignment(horizontal="center")

        adjust_columns_width(ws, start_row=4, min_width=12)
        wb.save(filepath)
        return filepath

    @staticmethod
    def generate_supply_excel(supply_data: dict, output_dir: str = "temp") -> str:
        """Priyomka (kirimlar) bo'yicha Excel hisoboti yaratadi"""
        os.makedirs(output_dir, exist_ok=True)
        filename = f"Kirimlar_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(output_dir, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Kirimlar"
        ws.views.sheetView[0].showGridLines = True

        ws.merge_cells("A1:F2")
        title_cell = ws["A1"]
        title_cell.value = "DE LUXURY — TOVAR KIRIMLARI (PRIYOMKA) HISOBOTI"
        title_cell.font = TITLE_FONT
        title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("A3:F3")
        sub_cell = ws["A3"]
        sub_cell.value = f"Shakllantirilgan sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        sub_cell.font = SUBTITLE_FONT
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")

        headers = ["№", "Hujjat raqami", "Sana va vaqt", "Kontragent (Yetkazib beruvchi)", "Pozitsiyalar soni", "Jami summa (so'm)"]

        row_num = 4
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        row_num = 5
        idx = 1
        for s in supply_data.get("rows", []):
            values = [
                idx, s.get("name", ""), s.get("moment", ""), s.get("agent", ""),
                s.get("positions_count", 0), s.get("sum", 0.0)
            ]
            for col_num, val in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.font = REGULAR_FONT
                cell.border = THIN_BORDER
                if col_num in [5, 6]:
                    cell.alignment = Alignment(horizontal="right")
                    if col_num == 6:
                        cell.number_format = "#,##0.00"
                elif col_num in [1, 2, 3]:
                    cell.alignment = Alignment(horizontal="center")
            row_num += 1
            idx += 1

        # Jami
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=5)
        ws.cell(row=row_num, column=1, value="JAMI KIRIM SUMMASI").font = TOTAL_FONT
        ws.cell(row=row_num, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=1).fill = TOTAL_FILL
        
        sum_cell = ws.cell(row=row_num, column=6, value=supply_data.get("total_sum", 0.0))
        sum_cell.font = TOTAL_FONT
        sum_cell.fill = TOTAL_FILL
        sum_cell.number_format = "#,##0.00"
        sum_cell.alignment = Alignment(horizontal="right")

        adjust_columns_width(ws, start_row=4, min_width=14)
        wb.save(filepath)
        return filepath

    @staticmethod
    def generate_supply_positions_excel(supply_positions_data: dict, supplier_name: str = "Barcha", output_dir: str = "temp") -> str:
        """Yetkazib beruvchi bo'yicha to'liq tovar pozitsiyalari Excel hisoboti"""
        os.makedirs(output_dir, exist_ok=True)
        clean_sup = "".join(c for c in supplier_name if c.isalnum() or c in (' ', '_', '-')).strip()
        filename = f"Kirim_{clean_sup}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(output_dir, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Kirim tovarlari"
        ws.views.sheetView[0].showGridLines = True

        # Sarlavhani A1:K2 da birlashtirib, chiroyli joylashtiramiz
        ws.merge_cells("A1:K2")
        title_cell = ws["A1"]
        title_cell.value = f"DE LUXURY — TOVAR KIRIMI (YETKAZIB BERUVCHI: {supplier_name.upper()})"
        title_cell.font = TITLE_FONT
        title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("A3:K3")
        sub_cell = ws["A3"]
        sub_cell.value = f"Shakllantirilgan sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        sub_cell.font = SUBTITLE_FONT
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")

        headers = [
            "№", "Hujjat №", "Sana va vaqt", "Tovar kodi", "Tovar nomi", "Guruhi", "Kategoriya", "Vid tovara",
            "Kirim soni (dona)", "Tan narxi (so'm)", "Jami summa (so'm)"
        ]

        row_num = 4
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        positions = supply_positions_data.get("positions", [])
        row_num = 5
        idx = 1
        total_qty = 0.0
        total_sum = 0.0

        for p in positions:
            qty = float(p.get("quantity", 0) or 0)
            price = float(p.get("price", 0) or 0)
            s_sum = float(p.get("sum", 0) or 0)

            total_qty += qty
            total_sum += s_sum

            values = [
                idx, p.get("supply_name", ""), p.get("moment", ""), p.get("code", ""), p.get("name", ""), p.get("folder", ""),
                p.get("category", ""), p.get("vid_tovara", ""),
                qty, price, s_sum
            ]

            for col_num, val in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.font = REGULAR_FONT
                cell.border = THIN_BORDER
                if col_num in [1, 2, 3, 4]:
                    cell.alignment = Alignment(horizontal="center")
                elif col_num in [5, 6, 7, 8]:
                    cell.alignment = Alignment(horizontal="left")
                elif col_num == 9:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0"
                elif col_num in [10, 11]:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0.00"

            row_num += 1
            idx += 1

        # Jami qatori
        total_cells = ["JAMI", "", "", "", "", "", "", "", total_qty, "", total_sum]
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=8)
        for col_num, val in enumerate(total_cells, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val if val != "" else None)
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = THIN_BORDER
            if col_num == 9:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0"
            elif col_num == 11:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.00"

        ws.cell(row=row_num, column=1).alignment = Alignment(horizontal="center")

        adjust_columns_width(ws, start_row=4, min_width=12)
        wb.save(filepath)
        return filepath
